"""Configuration partagée des tests pytest.

Pas de base de test isolée (S5 J43) : les tests d'intégration HTTP tapent
directement sur la base Postgres de dev, comme test_auth.py le fait déjà.
Chaque test utilise des données uniques (email via uuid4) pour ne jamais
entrer en collision avec un autre test ou une exécution précédente.
"""
import asyncio
import io
import uuid
from collections.abc import AsyncGenerator, Awaitable, Callable
from unittest.mock import AsyncMock, patch

import pytest
import pytest_asyncio
from httpx import ASGITransport, AsyncClient

from app.main import app


@pytest.fixture(scope="session")
def event_loop():
    """Une boucle d'événement par session de tests."""
    loop = asyncio.new_event_loop()
    yield loop
    loop.close()


@pytest_asyncio.fixture
async def client() -> AsyncGenerator[AsyncClient, None]:
    """Client HTTP async pour appeler l'API en tests."""
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        yield ac


# ============================================================
# Authentification — utilisateurs de test (S5 J43)
# ============================================================

DEFAULT_TEST_PASSWORD = "SecurePass123!"


@pytest.fixture
def fake():
    """Générateur de données factices en français."""
    from faker import Faker

    return Faker("fr_FR")


@pytest_asyncio.fixture
async def registered_user(client: AsyncClient) -> dict:
    """Inscrit un utilisateur unique (email aléatoire) et retourne ses
    identifiants en clair (utile pour re-login dans le test lui-même)."""
    email = f"test_{uuid.uuid4().hex[:10]}@example.com"
    response = await client.post(
        "/auth/register",
        json={"email": email, "password": DEFAULT_TEST_PASSWORD},
    )
    assert response.status_code == 201, response.text
    return {"email": email, "password": DEFAULT_TEST_PASSWORD, "id": response.json()["id"]}


@pytest_asyncio.fixture
async def auth_headers(client: AsyncClient, registered_user: dict) -> dict:
    """Headers Authorization Bearer pour un utilisateur fraîchement inscrit."""
    response = await client.post(
        "/auth/login",
        data={"username": registered_user["email"], "password": registered_user["password"]},
    )
    assert response.status_code == 200, response.text
    token = response.json()["access_token"]
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture
def register_second_user(
    client: AsyncClient,
) -> Callable[[], Awaitable[dict]]:
    """Factory fixture : cree un 2e utilisateur independant a la demande.

    Utile pour les tests d'ownership (ex: verifier qu'un user ne voit pas
    les donnees d'un autre). Usage : `other = await register_second_user()`.
    """

    async def _create() -> dict:
        email = f"test_{uuid.uuid4().hex[:10]}@example.com"
        await client.post(
            "/auth/register",
            json={"email": email, "password": DEFAULT_TEST_PASSWORD},
        )
        login = await client.post(
            "/auth/login",
            data={"username": email, "password": DEFAULT_TEST_PASSWORD},
        )
        token = login.json()["access_token"]
        return {"email": email, "headers": {"Authorization": f"Bearer {token}"}}

    return _create


# ============================================================
# Fichiers de test (PDF/Word/PowerPoint/CSV) — generes en memoire,
# jamais ecrits sur disque ni commites dans le repo (S5 J43)
# ============================================================

@pytest.fixture(scope="session")
def mini_pdf_bytes() -> bytes:
    """PDF de 2 pages avec texte simple, pour tester l'extraction PDF."""
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.drawString(100, 750, "Ceci est un document PDF de test.")
    c.drawString(100, 730, "Il contient deux pages.")
    c.showPage()
    c.drawString(100, 750, "Voici la deuxieme page.")
    c.save()
    return buf.getvalue()


@pytest.fixture(scope="session")
def mini_docx_bytes() -> bytes:
    """Word avec 3 paragraphes, pour tester l'extraction docx."""
    from docx import Document as DocxDocument

    buf = io.BytesIO()
    doc = DocxDocument()
    doc.add_paragraph("Premier paragraphe de test.")
    doc.add_paragraph(
        "Deuxieme paragraphe avec plus de contenu pour valider l'extraction."
    )
    doc.add_paragraph("Troisieme paragraphe pour terminer.")
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture(scope="session")
def mini_pptx_bytes() -> bytes:
    """PowerPoint avec 2 slides, pour tester l'extraction pptx."""
    from pptx import Presentation

    buf = io.BytesIO()
    prs = Presentation()
    slide1 = prs.slides.add_slide(prs.slide_layouts[0])
    slide1.shapes.title.text = "Slide 1"
    slide2 = prs.slides.add_slide(prs.slide_layouts[0])
    slide2.shapes.title.text = "Slide 2"
    prs.save(buf)
    return buf.getvalue()


@pytest.fixture(scope="session")
def mini_txt_bytes() -> bytes:
    """Texte brut UTF-8, pour tester l'extraction .txt."""
    return (
        "Ceci est un document texte de test.\n\n"
        "Il contient plusieurs paragraphes pour valider l'extraction "
        "et le decoupage en blocs.\n\n"
        "Troisieme paragraphe pour terminer proprement."
    ).encode("utf-8")


@pytest.fixture(scope="session")
def mini_md_bytes() -> bytes:
    """Markdown avec 2 titres, pour tester l'extraction .md avec sections."""
    return (
        "# Introduction\n\n"
        "Premiere section du document markdown de test.\n\n"
        "## Details\n\n"
        "Deuxieme section avec un sous-titre, pour verifier le decoupage "
        "sur les headers."
    ).encode("utf-8")


@pytest.fixture(scope="session")
def mini_csv_bytes() -> bytes:
    """CSV de 10 lignes avec colonnes typees (id/name/value/date)."""
    lines = ["id,name,value,date"]
    for i in range(10):
        lines.append(f"{i},Item {i},{i * 100},2026-01-{i + 1:02d}")
    return ("\n".join(lines) + "\n").encode("utf-8")


# ============================================================
# Mock des appels OpenAI externes (embeddings) — S5 J43
# ============================================================

async def _fake_embed(texts: list[str], model: str | None = None) -> list[list[float]]:
    return [[0.01] * 1536 for _ in texts]


@pytest.fixture
def mock_embeddings():
    """Mocke OpenAIProvider.embed (utilise par l'ingestion de documents et le
    retrieval RAG) pour ne jamais appeler la vraie API OpenAI en test."""
    with patch(
        "app.ai.providers.openai.OpenAIProvider.embed",
        new=AsyncMock(side_effect=_fake_embed),
    ):
        yield


@pytest_asyncio.fixture
async def db_session():
    """Session SQLAlchemy directe sur la meme DB que l'API (S5 J43).

    Reservee aux cas ou il n'existe pas de surface HTTP adaptee pour tester
    une logique service isolement (ex: retrieval RAG, qui a besoin d'inserer
    des chunks avec un embedding pgvector precisement controle). Pas de
    rollback : les lignes creees restent, comme partout ailleurs dans cette
    suite (utilisateurs/documents jetables, jamais vus par le vrai compte).
    """
    from app.db.session import AsyncSessionLocal

    async with AsyncSessionLocal() as session:
        yield session