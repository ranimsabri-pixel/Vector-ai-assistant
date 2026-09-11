"""Extraction de texte depuis un fichier .docx (S5 J32).

Une "section" est délimitée par les paragraphes de style Heading (ou Title).
S'il n'y a aucun titre dans le document, on découpe en groupes de
FALLBACK_GROUP_SIZE paragraphes pour rester proche du grain "page" du PDF.

Les tableaux sont convertis en texte tabulé simple (\t entre colonnes).
Les images sont ignorées (paragraph.text ne contient jamais de contenu image).
"""
from __future__ import annotations

from pathlib import Path

from docx import Document as DocxDocument
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

# Nombre de paragraphes/tableaux par section quand le document n'a aucun titre
FALLBACK_GROUP_SIZE = 10


def _iter_block_items(doc: DocxDocument):
    """Itère paragraphes et tableaux dans l'ordre du document.

    Recette standard python-docx : document.paragraphs et document.tables
    sont deux listes séparées qui ne préservent pas l'ordre d'apparition
    réel dans le fichier, il faut donc parcourir le XML du body directement.
    """
    for child in doc.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, doc)
        elif isinstance(child, CT_Tbl):
            yield Table(child, doc)


def _is_heading(paragraph: Paragraph) -> bool:
    style_name = (paragraph.style.name if paragraph.style else "") or ""
    return style_name.startswith("Heading") or style_name == "Title"


def _table_to_text(table: Table) -> str:
    """Convertit un tableau en texte tabule (\t entre colonnes).

    Piege python-docx : quand des cellules sont fusionnees horizontalement,
    row.cells renvoie le MEME objet cellule (meme _tc XML) une fois par
    colonne de grille couverte par la fusion -- sans dedup, une fusion sur
    5 colonnes duplique son texte 5 fois dans la ligne. On detecte les
    doublons par identite de _tc (pas par egalite de texte, pour ne pas
    ecraser des cellules distinctes qui contiendraient le meme texte).
    """
    lines = []
    for row in table.rows:
        seen_tc = None
        cells = []
        for cell in row.cells:
            if cell._tc is seen_tc:
                continue  # continuation d'une fusion horizontale, pas une vraie cellule
            seen_tc = cell._tc
            cells.append(cell.text.strip())
        lines.append("\t".join(cells))
    return "\n".join(lines)


def extract_docx_text(file_path: str) -> list[dict]:
    """Extrait le texte d'un fichier .docx, section par section.

    Retourne une liste de dicts :
        [{ "section_number": 1, "section_title": "Introduction", "content": "..." }, ...]
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"Fichier DOCX introuvable : {file_path}")

    doc = DocxDocument(str(path))
    blocks = list(_iter_block_items(doc))

    has_headings = any(
        isinstance(b, Paragraph) and _is_heading(b) and b.text.strip()
        for b in blocks
    )

    sections: list[dict] = []
    current_title: str | None = None
    buffer: list[str] = []
    section_number = 0
    group_count = 0

    def flush() -> None:
        nonlocal section_number, buffer
        content = "\n".join(buffer).strip()
        if content:
            section_number += 1
            sections.append({
                "section_number": section_number,
                "section_title": current_title,
                "content": content,
            })
        buffer = []

    for block in blocks:
        if isinstance(block, Paragraph):
            if has_headings and _is_heading(block):
                if block.text.strip():
                    flush()
                    current_title = block.text.strip()
                    group_count = 0
                continue

            text = block.text.strip()
            if not text:
                continue
            buffer.append(text)
            group_count += 1
        else:
            table_text = _table_to_text(block)
            if not table_text.strip():
                continue
            buffer.append(table_text)
            group_count += 1

        # Pas de titres dans le document : on regroupe par lots fixes
        if not has_headings and group_count >= FALLBACK_GROUP_SIZE:
            flush()
            group_count = 0

    flush()

    if not sections:
        raise ValueError(
            "Aucun texte extractible dans ce document Word (fichier vide ou "
            "contenu uniquement composé d'images)."
        )

    return sections
