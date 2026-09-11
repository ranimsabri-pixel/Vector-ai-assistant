"""Extraction de texte depuis un fichier .pptx (S5 J35).

Une "section" correspond a une slide. Le titre est le placeholder titre
de la slide (slide.shapes.title) quand il existe. Le contenu regroupe les
autres zones de texte, les tableaux (convertis en texte tabule) et les
notes du presentateur. Les images sont ignorees.
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.table import Table


def _table_to_text(table: Table) -> str:
    lines = []
    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells]
        lines.append("\t".join(cells))
    return "\n".join(lines)


def extract_pptx_text(file_path: str) -> list[dict]:
    """Extrait le texte d'une presentation .pptx, slide par slide.

    Retourne une liste de dicts :
        [{ "section_number": 1, "section_title": "Titre slide", "content": "..." }, ...]
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"Fichier PPTX introuvable : {file_path}")

    prs = Presentation(str(path))
    sections: list[dict] = []

    for slide_num, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        title_text: str | None = None
        if title_shape is not None and title_shape.has_text_frame:
            title_text = title_shape.text_frame.text.strip() or None

        title_shape_id = title_shape.shape_id if title_shape is not None else None

        content_parts: list[str] = []
        for shape in slide.shapes:
            # slide.shapes.title et l'iteration slide.shapes renvoient des
            # objets Python distincts pour le meme element XML : comparer
            # par shape_id, pas par identite d'objet ("is" ne marche pas ici).
            if title_shape_id is not None and shape.shape_id == title_shape_id:
                continue  # deja capture comme titre, evite la duplication

            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        content_parts.append(text)
            elif shape.has_table:
                table_text = _table_to_text(shape.table)
                if table_text.strip():
                    content_parts.append(table_text)

        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                content_parts.append(f"[Notes] {notes}")

        content = "\n".join(content_parts).strip()
        if not content and not title_text:
            continue  # slide vide (ou uniquement des images), on saute

        sections.append({
            "section_number": slide_num,
            "section_title": title_text,
            "content": content or title_text or "",
        })

    if not sections:
        raise ValueError(
            "Aucun texte extractible dans cette présentation (slides vides "
            "ou composées uniquement d'images)."
        )

    return sections
